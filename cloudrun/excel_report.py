"""Deterministic workbook for a single Google Cloud inventory identity."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from classification import classify, counts

HEADERS = {
    "Overview": ["Project ID", "Primary resources", "Configuration", "System", "Other", "Enabled APIs"],
    "Resources": ["Project ID", "Resource", "Asset type", "Category", "Layer", "Confidence", "Location", "State"],
    "Networking": ["Project ID", "Rule", "Direction", "Sources", "Protocols / ports", "Enabled"],
    "IAM": ["Project ID", "Service account", "Disabled", "IAM binding count", "IAM roles"],
    "Enabled APIs": ["Project ID", "API", "Title"],
    "Security Review": ["Project ID", "Category", "Resource", "Observation"],
    "Raw Inventory": ["Project ID", "Name", "Asset type", "Location", "State"],
}

def workbook_bytes(projects, deep_results=None):
    from observations import observations
    deep_by_project = {x["projectId"]: x for x in (deep_results or [])}
    wb = Workbook()
    wb.remove(wb.active)
    sheets = {name: wb.create_sheet(name) for name in HEADERS}
    for name, headers in HEADERS.items():
        sheets[name].append(headers)
    for p in sorted(projects, key=lambda x: x["projectId"]):
        pid = p["projectId"]
        resources = sorted(p.get("resources", []), key=lambda r: (r.get("assetType") or "", r.get("name") or ""))
        c = counts(resources)
        sheets["Overview"].append([pid, c["primary"], c["configuration"], c["system"], c["unclassified"] + c["api"], len(p.get("apis", []))])
        for r in resources:
            kind = classify(r)
            row = [pid, r.get("name"), r.get("assetType"), kind["category"], kind["layer"], kind["confidence"], r.get("location"), r.get("state")]
            sheets["Resources"].append(row)
            sheets["Raw Inventory"].append([pid, r.get("name"), r.get("assetType"), r.get("location"), r.get("state")])
        for api in sorted(p.get("apis", []), key=lambda a: a.get("name") or ""):
            sheets["Enabled APIs"].append([pid, api.get("name"), api.get("title")])
        deep = deep_by_project.get(pid, {})
        services = deep.get("services", {})
        for rule in services.get("firewall", []) if isinstance(services.get("firewall"), list) else []:
            allowed = "; ".join(str(x.get("protocol", "")) + ":" + ",".join(x.get("ports", [])) for x in rule.get("allowed", []))
            sheets["Networking"].append([pid, rule.get("name"), rule.get("direction"), ", ".join(rule.get("sourceRanges", [])), allowed, not rule.get("disabled", False)])
        policy = services.get("iamPolicy", {})
        if not isinstance(policy, dict) or policy.get("status") == "not_enabled":
            policy = {}
        accounts = services.get("serviceAccounts", [])
        if isinstance(accounts, list):
            for account in accounts:
                sheets["IAM"].append([pid, account.get("email"), account.get("disabled"), policy.get("bindingsCount"), ", ".join(policy.get("roles", []))])
        for finding in observations(deep):
            sheets["Security Review"].append([pid, finding.get("category"), finding.get("resource"), finding.get("detail")])
    for name, ws in sheets.items():
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 28
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="203B5C")
            cell.alignment = Alignment(vertical="center")
        for col in ws.columns:
            letter = col[0].column_letter
            maxlen = max(len(str(cell.value or "")) for cell in list(col)[:150])
            ws.column_dimensions[letter].width = min(max(15, maxlen + 2), 55)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
