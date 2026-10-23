import unittest
from io import BytesIO
from openpyxl import load_workbook
from excel_report import workbook_bytes

class ExcelTests(unittest.TestCase):
    def test_workbook_sheets_and_counts(self):
        projects = [{"projectId": "test", "resources": [
            {"name": "bucket-a", "assetType": "storage.googleapis.com/Bucket"},
            {"name": "other", "assetType": "unknown.googleapis.com/Other"},
            {"name": "api-asset", "assetType": "serviceusage.googleapis.com/Service"}
        ], "apis": [{"name": "storage.googleapis.com", "title": "Storage"}]}]
        data = workbook_bytes(projects)
        wb = load_workbook(BytesIO(data), read_only=True)
        self.assertEqual(len(wb.sheetnames), 7)
        self.assertEqual(wb["Overview"]["B2"].value, 1)
        self.assertEqual(wb["Overview"]["E2"].value, 1)
        self.assertEqual(wb["Overview"]["F2"].value, 1)
        self.assertEqual(wb["Overview"]["E1"].value, "Unclassified")
        self.assertEqual(wb["Resources"]["A2"].value, "test")

if __name__ == "__main__":
    unittest.main()
